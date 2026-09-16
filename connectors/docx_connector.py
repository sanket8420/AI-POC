"""
DocxConnector
=============
Implements BaseConnector for Word documents (.docx).

Unlike PDFs, docx tables are stored as REAL structured objects (rows
and cells) — no position-based reconstruction hack needed. This
connector reads native tables directly, and runs the shared prose
pipeline (page_processor) on the surrounding paragraph text for
key-value/price-anchor detection.

Connect  -> open with python-docx
Extract  -> pull all paragraph text + all native tables
Validate -> ensure the document has SOME usable content
Transform-> classify+extract the prose; convert native tables directly
            (no table_parser needed — python-docx already gives clean
            rows/cells)
"""

import os
from typing import Any, Dict, List

import docx

from connectors.base_connector import BaseConnector
from document_intelligence.page_processor import process_text
from document_intelligence.table_parser import normalize_key


class DocxConnector(BaseConnector):
    def __init__(self, file_path: str, min_text_length: int = 5):
        super().__init__(source_name="docx")
        self.file_path = file_path
        self.min_text_length = min_text_length

    def connect(self) -> None:
        if not os.path.isfile(self.file_path):
            raise FileNotFoundError(f"DOCX file not found: {self.file_path}")
        self._connection = docx.Document(self.file_path)

    def extract(self) -> List[Dict[str, Any]]:
        doc = self._connection

        paragraph_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        native_tables = []
        for table in doc.tables:
            rows_data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if len(rows_data) < 2:
                continue  # need at least a header + one data row
            header = [normalize_key(h) for h in rows_data[0]]
            for row in rows_data[1:]:
                native_tables.append(dict(zip(header, row)))

        return [{"raw_text": paragraph_text, "native_table_rows": native_tables}]

    def validate(self, raw_records: List[Dict[str, Any]]):
        valid, failed = [], []
        for record in raw_records:
            has_text = len((record.get("raw_text") or "").strip()) >= self.min_text_length
            has_table = bool(record.get("native_table_rows"))
            if has_text or has_table:
                valid.append(record)
            else:
                failed_record = dict(record)
                failed_record["_error"] = "No usable paragraph text or tables found in this document"
                failed.append(failed_record)
        return valid, failed

    def transform(self, valid_records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        transformed = []
        for i, record in enumerate(valid_records):
            processed = process_text(record["raw_text"]) if record["raw_text"].strip() else {
                "cleaned_text": "", "document_type_guess": "unknown",
                "classification_confidence": 0.0, "extraction_method": "raw_text",
                "item_count": 0, "items": [],
            }

            native_rows = record.get("native_table_rows", [])
            if native_rows:
                # Native docx tables are always reliable — merge them in
                # alongside whatever the prose pipeline found.
                processed["items"] = processed["items"] + native_rows
                processed["item_count"] = len(processed["items"])
                processed["extraction_method"] = (
                    processed["extraction_method"] + "+native_docx_table"
                    if processed["extraction_method"] != "raw_text"
                    else "native_docx_table"
                )

            transformed.append({
                "record_id": f"docx-{i}",
                "source_type": "docx",
                "content": processed,
                "metadata": {"origin_file": os.path.basename(self.file_path)},
            })
        return transformed
